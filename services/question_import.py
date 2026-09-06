import csv
import io
from openpyxl import load_workbook

REQUIRED = ['question_text','option_a','option_b','option_c','option_d','correct_answer']
OPTIONAL = ['marks','topic','difficulty','explanation']
ALIASES = {
    'question': 'question_text', 'question text': 'question_text', 'question_text': 'question_text',
    'option a': 'option_a', 'option_a': 'option_a', 'a': 'option_a',
    'option b': 'option_b', 'option_b': 'option_b', 'b': 'option_b',
    'option c': 'option_c', 'option_c': 'option_c', 'c': 'option_c',
    'option d': 'option_d', 'option_d': 'option_d', 'd': 'option_d',
    'correct answer': 'correct_answer', 'correct_answer': 'correct_answer', 'answer': 'correct_answer', 'correct': 'correct_answer',
    'marks': 'marks', 'mark': 'marks', 'topic': 'topic', 'difficulty': 'difficulty', 'explanation': 'explanation'
}

def _norm(s):
    return ' '.join(str(s or '').strip().lower().replace('_', ' ').split())

def _map_headers(headers):
    mapped = {}
    for i, h in enumerate(headers):
        key = ALIASES.get(_norm(h))
        if key and key not in mapped:
            mapped[key] = i
    return mapped

def read_question_file(file_storage):
    filename = (file_storage.filename or '').lower()
    if filename.endswith('.xlsx'):
        wb = load_workbook(file_storage.stream, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try: headers = next(rows)
        except StopIteration: return [], ['The Excel file is empty.']
        rows = list(rows)
    elif filename.endswith('.csv'):
        raw = file_storage.read()
        text = raw.decode('utf-8-sig', errors='replace')
        parsed = list(csv.reader(io.StringIO(text)))
        if not parsed: return [], ['The CSV file is empty.']
        headers, rows = parsed[0], parsed[1:]
    else:
        return [], ['Only .xlsx and .csv files are supported.']

    mapped = _map_headers(headers)
    missing = [h for h in REQUIRED if h not in mapped]
    if missing:
        return [], ['Missing required columns: ' + ', '.join(missing)]

    records, errors = [], []
    for row_no, row in enumerate(rows, start=2):
        def val(key):
            idx = mapped.get(key)
            return str(row[idx]).strip() if idx is not None and idx < len(row) and row[idx] is not None else ''
        if not any(str(x or '').strip() for x in row):
            continue
        correct = val('correct_answer').upper()
        if correct not in {'A','B','C','D'}:
            errors.append(f'Row {row_no}: correct_answer must be A, B, C, or D.')
            continue
        if any(not val(k) for k in REQUIRED if k != 'correct_answer'):
            errors.append(f'Row {row_no}: question and all four options are required.')
            continue
        try: marks = float(val('marks') or 1)
        except ValueError:
            errors.append(f'Row {row_no}: marks must be a number.')
            continue
        records.append({
            'question_text': val('question_text'), 'option_a': val('option_a'), 'option_b': val('option_b'),
            'option_c': val('option_c'), 'option_d': val('option_d'), 'correct_answer': correct,
            'marks': marks, 'topic': val('topic') or 'General', 'difficulty': val('difficulty') or 'Medium',
            'explanation': val('explanation')
        })
    return records, errors
