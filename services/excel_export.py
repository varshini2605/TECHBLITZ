import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from services.scoring import get_ranked_results

def generate_results_excel(test_id=None):
    """
    Generates an Excel workbook with results styled professionally with
    headers, formatting, rankings, and proctoring stats.
    Returns bytes buffer of TECHBLITZ_RESULTS.xlsx
    """
    ranked_candidates = get_ranked_results(test_id=test_id)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TECHBLITZ Results"
    
    # Enable gridlines
    ws.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws.merge_cells('A1:P1')
    title_cell = ws['A1']
    title_cell.value = "TECHBLITZ — Online Assessment Platform | Official Examination Results"
    title_cell.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1E1B4B', end_color='1E1B4B', fill_type='solid') # Deep Indigo
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40
    
    # Subtitle
    ws.merge_cells('A2:P2')
    sub_cell = ws['A2']
    sub_cell.value = f"Rankings sorted by Highest Score → Lowest Score (Tie-breaker: Earlier submission time) | Total Candidates: {len(ranked_candidates)}"
    sub_cell.font = Font(name='Calibri', size=10, italic=True, color='4B5563')
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20
    
    # Headers
    headers = [
        "Rank", "Candidate ID", "Hall Ticket Number", "Name", "Team ID",
        "Email", "Test ID", "Score", "Total Marks", "Percentage (%)",
        "Violations", "Violation Types", "Status", "Start Time",
        "Submission Time", "Time Taken"
    ]
    
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid') # Indigo primary
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    border_thin = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    ws.row_dimensions[4].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = cell_border
        
    # Data rows
    row_alt_fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
    
    for r_idx, c in enumerate(ranked_candidates, start=5):
        ws.row_dimensions[r_idx].height = 22
        fill = row_alt_fill if r_idx % 2 == 0 else PatternFill(fill_type=None)
        
        row_values = [
            c['rank'],
            c['candidate_id'],
            c['hall_ticket_number'],
            c['name'],
            c['team_id'],
            c['email'],
            c['test_id'],
            c['score'],
            c['total_marks'],
            f"{c['percentage']:.1f}%",
            c['violations_count'],
            c['violation_types'],
            c['status'],
            c['start_time'],
            c['submission_time'],
            c['time_taken']
        ]
        
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10)
            cell.border = cell_border
            if fill.fill_type:
                cell.fill = fill
                
            # Alignment rules
            if col_idx in [1, 5, 8, 9, 10, 11, 13, 16]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in [2, 3, 4, 6, 7, 12, 14, 15]:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
    from openpyxl.utils import get_column_letter
    # Auto-fit columns
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            if cell.row in [1, 2]: # skip title merges
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
