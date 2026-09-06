import os
import openpyxl
import csv

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), '..', 'sample_data')
os.makedirs(SAMPLE_DIR, exist_ok=True)

# 1. Standard registrations_sample.xlsx
wb1 = openpyxl.Workbook()
ws1 = wb1.active
ws1.title = "Registrations"
ws1.append(["Hall Ticket Number", "Name", "Team ID", "Email"])
sample_candidates = [
    ["JITS23A0101", "Aakash Sharma", "TEAM-01", "aakash@college.edu"],
    ["JITS23A0102", "Bhavana Patel", "TEAM-01", "bhavana@college.edu"],
    ["JITS23A0103", "Chirag Sen", "TEAM-02", "chirag@college.edu"],
    ["JITS23A0104", "Deepika Roy", "TEAM-02", "deepika@college.edu"],
    ["JITS23A0105", "Eshwar Varma", "TEAM-03", "eshwar@college.edu"],
    ["JITS23A0106", "Farhan Khan", "TEAM-03", "farhan@college.edu"],
    ["JITS23A0107", "Gautami Nair", "TEAM-04", "gautami@college.edu"],
    ["JITS23A0108", "Harshavardhan", "TEAM-04", "harsha@college.edu"],
    ["JITS23A0109", "Ishita Gupta", "TEAM-05", "ishita@college.edu"],
    ["JITS23A0110", "Jaspreet Singh", "TEAM-05", "jaspreet@college.edu"]
]
for row in sample_candidates:
    ws1.append(row)
wb1.save(os.path.join(SAMPLE_DIR, "registrations_sample.xlsx"))

# 2. Alternate headers registrations_alternate_headers.xlsx
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Form Responses"
ws2.append(["HT No", "Student Name", "Team Number", "Email Address"])
alt_candidates = [
    ["JITS23B0201", "Manish Joshi", "ALPHA-01", "manish@college.edu"],
    ["JITS23B0202", "Nandini Das", "ALPHA-01", "nandini@college.edu"],
    ["JITS23B0203", "Omkar Deshmukh", "BETA-02", "omkar@college.edu"],
    ["JITS23B0204", "Pooja Hegde", "BETA-02", "pooja@college.edu"]
]
for row in alt_candidates:
    ws2.append(row)
wb2.save(os.path.join(SAMPLE_DIR, "registrations_alternate_headers.xlsx"))

# 3. Duplicate handling registrations_with_duplicates.xlsx
wb3 = openpyxl.Workbook()
ws3 = wb3.active
ws3.title = "Duplicates Test"
ws3.append(["Hall Ticket", "Name", "Team ID", "Email"])
dup_candidates = [
    ["TB001", "Rahul Updated", "TEAM01", "rahul.updated@college.edu"], # Existing in seed
    ["TB002", "Priya Updated", "TEAM01", "priya.updated@college.edu"], # Existing in seed
    ["JITS23C0301", "Rohan Mehta", "GAMMA-01", "rohan@college.edu"],
    ["JITS23C0301", "Rohan Duplicate", "GAMMA-01", "rohan.dup@college.edu"], # duplicate within file
    ["JITS23C0302", "", "GAMMA-02", "missing.name@college.edu"], # missing name
    ["JITS23C0303", "Sneha Roy", "", "missing.team@college.edu"] # missing team
]
for row in dup_candidates:
    ws3.append(row)
wb3.save(os.path.join(SAMPLE_DIR, "registrations_with_duplicates.xlsx"))

# 4. Sample CSV
with open(os.path.join(SAMPLE_DIR, "registrations_sample.csv"), "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(["Hall Ticket Number", "Name", "Team ID", "Email"])
    for row in sample_candidates[:5]:
        writer.writerow(row)

print("Generated sample registration files successfully in sample_data/")
